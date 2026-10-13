# -*- coding: utf-8 -*-
# OSInt Graph - esplorazione grafica di dati OSINT per casi investigativi
# Copyright (C) 2025-2026 Andrea Cumini <andrea@osintinfo.net> - www.osintinfo.net
# SPDX-License-Identifier: GPL-3.0-or-later
#
# This program is free software: you can redistribute it and/or modify it
# under the terms of the GNU General Public License as published by the Free
# Software Foundation, either version 3 of the License, or (at your option)
# any later version, with the additional terms in the file NOTICE.
# See the files LICENSE and NOTICE for details. Distributed WITHOUT ANY WARRANTY.

"""Indexing of a folder of sources (documents) and full-text search.

Supported formats for text extraction:
    .txt .md .csv .log    read directly
    .docx                 python-docx
    .doc                  best-effort (readable strings from the binary)
    .xlsx                 openpyxl
    .xls                  xlrd
    .pdf                  pypdf (text PDFs)
    .png .jpg .jpeg ...   OCR with Tesseract, ONLY if installed

The caller saves the index (in the app database) so it survives restarts.
Search is by substring (partial) or by whole word (total), case- and
accent-insensitive, and returns context previews.
"""
# "Mendacem memorem esse oportet." (Quintilian, Institutio Oratoria)
# A liar needs a good memory.
import os
import re
import io
import json
import time
import hashlib
import logging
import itertools
import unicodedata

TEXT_EXT = {".txt", ".md", ".csv", ".log", ".text", ".json"}
DOCX_EXT = {".docx"}
DOC_EXT = {".doc"}
XLSX_EXT = {".xlsx", ".xlsm"}
XLS_EXT = {".xls"}
PDF_EXT = {".pdf"}
IMG_EXT = {".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp", ".tif", ".tiff"}
SUPPORTED = TEXT_EXT | DOCX_EXT | DOC_EXT | XLSX_EXT | XLS_EXT | PDF_EXT | IMG_EXT

MAX_TEXT = 800_000          # maximum text kept per file (characters)
MAX_FILE_MB = 60            # files larger than this are skipped


# ---------------------------------------------------------------- OCR ------
# Tesseract can live in a LOCAL folder (downloaded with setup_tesseract.py),
# without installing it system-wide. It is looked for there before the PATH.
def _local_tesseract_candidates(base_dir):
    exe = "tesseract.exe" if os.name == "nt" else "tesseract"
    dirs = []
    if os.environ.get("TESSERACT_DIR"):
        dirs.append(os.environ["TESSERACT_DIR"])
    if base_dir:
        dirs.append(os.path.join(base_dir, "tools", "tesseract"))
    out = []
    for d in dirs:
        out.append(os.path.join(d, exe))
        out.append(os.path.join(d, "bin", exe))       # layout with a bin/ subfolder
    return out


TESSDATA_DIR = ""          # local tessdata folder, if found


OCR_REASON = ""     # why OCR is missing: "module" (pytesseract), "binary" (tesseract) or ""
OCR_LANG = ""       # languages passed to Tesseract ("ita+eng", eng last): TESSERACT_LANG or all installed ones


def _ocr_languages():
    """Languages to pass to Tesseract: the TESSERACT_LANG variable if set,
    otherwise every installed language pack (except 'osd'), joined with '+'.
    With no packs or on error, Tesseract decides (eng)."""
    env = os.environ.get("TESSERACT_LANG", "").strip()
    if env:
        return env
    try:
        import pytesseract
        langs = [l for l in pytesseract.get_languages(config="") if l and l != "osd"]
    except Exception:
        return ""
    # the FIRST language drives recognition: eng last, so accents and
    # alphabets of the other languages are not flattened into English
    langs.sort(key=lambda l: (l == "eng", l))
    return "+".join(langs)


def configure_ocr(base_dir=None):
    """Configure pytesseract with a local Tesseract if present, otherwise the
    system one. Updates and returns OCR_AVAILABLE (and OCR_REASON)."""
    global OCR_AVAILABLE, TESSDATA_DIR, OCR_REASON, OCR_LANG
    try:
        import pytesseract
    except Exception:
        OCR_AVAILABLE, OCR_REASON = False, "module"
        return False
    # 1) local binary (tools/tesseract folder)
    for cand in _local_tesseract_candidates(base_dir):
        if os.path.isfile(cand):
            pytesseract.pytesseract.tesseract_cmd = cand
            # the language data sits alongside, in tessdata/. Tesseract 5.x wants
            # TESSDATA_PREFIX pointing DIRECTLY at that folder.
            tdata = os.path.join(os.path.dirname(cand), "tessdata")
            if os.path.isdir(tdata):
                TESSDATA_DIR = tdata
                os.environ["TESSDATA_PREFIX"] = tdata
            break
    # 2) check that it works (local or system)
    try:
        pytesseract.get_tesseract_version()
        OCR_AVAILABLE, OCR_REASON = True, ""
        OCR_LANG = _ocr_languages()
    except Exception:
        OCR_AVAILABLE, OCR_REASON = False, "binary"
    return OCR_AVAILABLE


OCR_AVAILABLE = configure_ocr()


# -------------------------------------------------------- text extractors ---
def _txt(path):
    for enc in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            with io.open(path, "r", encoding=enc) as f:
                return f.read()
        except Exception:
            continue
    return ""


def _docx(path):
    import docx
    d = docx.Document(path)
    parts = [p.text for p in d.paragraphs]
    for tbl in d.tables:                        # text in tables too
        for row in tbl.rows:
            parts.append("\t".join(c.text for c in row.cells))
    return "\n".join(parts)


def _doc(path):
    """Old .doc format: extracts readable character runs from the binary.
    Crude, but it captures the visible text without external dependencies."""
    try:
        raw = io.open(path, "rb").read()
    except Exception:
        return ""
    txt = raw.decode("latin-1", "ignore")
    tokens = re.findall(r"[ -~ -ɏ]{4,}", txt)   # ASCII + extended Latin
    return "\n".join(t.strip() for t in tokens if len(t.strip()) >= 4)


def _xlsx(path):
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    out = []
    for ws in wb.worksheets:
        out.append(f"# {ws.title}")
        for row in ws.iter_rows(values_only=True):
            cells = [str(c) for c in row if c not in (None, "")]
            if cells:
                out.append("\t".join(cells))
    wb.close()
    return "\n".join(out)


def _xls(path):
    import xlrd
    book = xlrd.open_workbook(path)
    out = []
    for sh in book.sheets():
        out.append(f"# {sh.name}")
        for r in range(sh.nrows):
            cells = [str(c) for c in sh.row_values(r) if str(c).strip()]
            if cells:
                out.append("\t".join(cells))
    return "\n".join(out)


def _pdf(path):
    from pypdf import PdfReader
    r = PdfReader(path)
    out = []
    for pg in r.pages:
        try:
            out.append(pg.extract_text() or "")
        except Exception:
            pass
        if sum(len(x) for x in out) > MAX_TEXT:
            break
    return "\n".join(out)


def _image(path):
    if not OCR_AVAILABLE:
        return ""
    try:
        import pytesseract
        from PIL import Image
        # the language data location comes from TESSDATA_PREFIX (set in
        # configure_ocr directly on the tessdata folder, as Tesseract 5 wants).
        # languages: all installed ones (or TESSERACT_LANG), so accents and
        # non-Latin alphabets are read; without language packs it stays eng.
        return pytesseract.image_to_string(Image.open(path), lang=OCR_LANG or None)
    except Exception as e:
        logging.warning(f"OCR failed on {path}: {e}")
        return ""


def extract_text(path, ext):
    try:
        if ext in TEXT_EXT:
            return _txt(path)
        if ext in DOCX_EXT:
            return _docx(path)
        if ext in DOC_EXT:
            return _doc(path)
        if ext in XLSX_EXT:
            return _xlsx(path)
        if ext in XLS_EXT:
            return _xls(path)
        if ext in PDF_EXT:
            return _pdf(path)
        if ext in IMG_EXT:
            return _image(path)
    except Exception as e:
        logging.warning(f"Extraction failed on {path}: {e}")
    return ""


# ---------------------------------------------------------- normalisation ---
def _norm_char(c):
    """Normalise ONE character staying 1:1 (one char -> one char): strips the
    accent and lowercases, without changing the length. Essential so that
    match offsets line up with the original text (highlight)."""
    d = unicodedata.normalize("NFKD", c)
    d = "".join(ch for ch in d if not unicodedata.combining(ch))
    return (d[0] if d else c).lower()


def norm(s):
    """Lowercase, accents removed, BUT with a 1:1 character mapping, so the
    position of each character is the same as in the original text."""
    return "".join(_norm_char(c) for c in str(s or ""))


def file_id(path):
    return hashlib.md5(os.path.abspath(path).encode("utf-8")).hexdigest()[:16]


# ------------------------------------------------------------------ scan ---
def _list_candidates(folder):
    """List of supported files (path, name, ext, size), skipping those too
    large. Used to know the TOTAL before extraction starts."""
    out, skipped = [], 0
    for root, dirs, names in os.walk(folder):
        dirs[:] = [d for d in dirs if not d.startswith(".")]
        for name in names:
            ext = os.path.splitext(name)[1].lower()
            if ext not in SUPPORTED:
                continue
            path = os.path.join(root, name)
            try:
                size = os.path.getsize(path)
            except OSError:
                continue
            if size > MAX_FILE_MB * 1024 * 1024:
                skipped += 1
                continue
            out.append((path, name, ext, size))
    return out, skipped


def _rel_to_base(folder, base_dir):
    """If 'folder' is inside base_dir, return the relative path (for
    PORTABILITY: moving the folder keeps the index valid). Otherwise None."""
    if not base_dir:
        return None
    try:
        base = os.path.abspath(base_dir)
        rel = os.path.relpath(folder, base)
    except ValueError:                      # different drives on Windows
        return None
    # inside base_dir if the relative path does not climb with '..' and is not absolute
    if rel == os.pardir or rel.startswith(os.pardir + os.sep) or os.path.isabs(rel):
        return None
    return rel.replace("\\", "/")


def effective_folder(index, base_dir):
    """CURRENT sources folder: if the index has a relative path (folder
    inside/next to the executable) it is resolved against base_dir NOW, so
    everything stays valid even after the folder has been moved."""
    rel = index.get("folder_rel")
    if rel and base_dir:
        return os.path.abspath(os.path.join(base_dir, rel.replace("/", os.sep)))
    return index.get("folder") or ""


SOURCES_FOLDER = "sources"
LEGACY_FOLDER = "fonti"      # name of the folder up to 1.7.x


def sources_folder(base_dir):
    """Documents folder: the SOURCES_DIR environment variable when set (the
    documents kept outside the program folder, like DATA_DIR for the cases),
    otherwise 'sources' next to the program. Installations made before 1.8.0
    have 'fonti': as long as 'sources' does not exist and 'fonti' does, the
    latter is still used, so the index and documents stay valid."""
    if os.environ.get("SOURCES_DIR"):
        return os.path.abspath(os.environ["SOURCES_DIR"])
    new = os.path.join(base_dir, SOURCES_FOLDER)
    old = os.path.join(base_dir, LEGACY_FOLDER)
    if not os.path.isdir(new) and os.path.isdir(old):
        return old
    return new


def folder_candidates(index, base_dir, default_folder=None):
    """Folders in which to look for the indexed documents TODAY, IN ORDER.

    The 'sources' folder next to the executable (or 'fonti' before 1.8.0, see
    sources_folder()) is ALWAYS checked first: it is the official location of
    the documents and the only one that stays valid when the program is moved
    from one PC to another.
      1) sources/<name of the scanned folder>   (e.g. sources/Tutti_dati_Lybia)
      2) sources/
      3) the path recorded in the index (absolute: valid only on the
         machine that did the scan)

    Without points 1-2 this happened: the index keeps the path of the
    scanned folder and, if that was OUTSIDE the program directory, the path
    is absolute. Moving executable + data/ + sources/ to another PC, search
    kept working (the text is inside the index) but opening files
    returned 404."""
    out = []

    def aggiungi(p):
        if not p:
            return
        p = os.path.abspath(p)
        if p not in out:
            out.append(p)

    fonti = default_folder or (sources_folder(base_dir) if base_dir else None)
    if fonti:
        originale = (index.get("folder") or "").replace("/", os.sep).rstrip("\\/")
        nome = os.path.basename(originale)
        if nome:
            aggiungi(os.path.join(fonti, nome))
        aggiungi(fonti)
        # last attempt inside the sources folder: the first-level subfolders.
        # Covers the case where the documents were put there under a name
        # different from that of the scanned folder.
        try:
            for voce in sorted(os.listdir(fonti))[:40]:
                sotto = os.path.join(fonti, voce)
                if os.path.isdir(sotto):
                    aggiungi(sotto)
        except OSError:
            pass
    aggiungi(effective_folder(index, base_dir))
    return out


def locate_file(index, rel, base_dir, default_folder=None):
    """Real path of an indexed file, looked up among the candidate folders.
    Returns (path, folder) or (None, None) if not found.
    The security check stays: the file must be INSIDE the folder."""
    rel = (rel or "").replace("/", os.sep)
    if not rel:
        return None, None
    for cartella in folder_candidates(index, base_dir, default_folder):
        percorso = os.path.abspath(os.path.join(cartella, rel))
        if not percorso.startswith(cartella + os.sep):
            continue                      # rel trying to escape: discarded
        if os.path.isfile(percorso):
            return percorso, cartella
    return None, None


def resolved_folder(index, base_dir, default_folder=None):
    """Folder the documents are ACTUALLY being read from: the first
    candidate that exists and contains at least one file of the index."""
    files = index.get("files") or []
    campione = [f.get("rel") for f in files[:25] if f.get("rel")]
    for cartella in folder_candidates(index, base_dir, default_folder):
        if not os.path.isdir(cartella):
            continue
        if not campione:
            return cartella
        for rel in campione:
            if os.path.isfile(os.path.join(cartella, rel.replace("/", os.sep))):
                return cartella
    return effective_folder(index, base_dir)


def scan_folder_stream(folder, index_path, base_dir=None):
    """Like scan_folder, but a GENERATOR that emits progress for every
    processed file: first {'phase':'start','total':N}, then for each file
    {'phase':'file','done':i,'total':N,'name':...}, finally {'phase':'done', ...}
    with the summary. Saves the index at the end. If the folder is inside
    base_dir, stores the RELATIVE path (portability)."""
    folder = os.path.abspath(folder)
    if not os.path.isdir(folder):
        raise FileNotFoundError(folder)

    folder_rel = _rel_to_base(folder, base_dir)
    candidates, skipped = _list_candidates(folder)
    total = len(candidates)
    yield {"phase": "start", "total": total, "skipped": skipped}

    files, by_ext = [], {}
    for i, (path, name, ext, size) in enumerate(candidates, 1):
        yield {"phase": "file", "done": i, "total": total, "name": name}
        text = extract_text(path, ext)[:MAX_TEXT]
        files.append({
            "id": file_id(os.path.relpath(path, folder)),   # id stable across moves
            "rel": os.path.relpath(path, folder).replace("\\", "/"),
            "name": name, "ext": ext,
            "size": size, "mtime": os.path.getmtime(path),
            "text": text, "ntext": norm(text), "chars": len(text),
            "ocr": ext in IMG_EXT and bool(text),
        })
        by_ext[ext] = by_ext.get(ext, 0) + 1

    index = {
        "folder": folder,             # absolute: for reference/display only
        "folder_rel": folder_rel,     # relative to base_dir if portable
        "scanned_at": time.strftime("%Y-%m-%d %H:%M:%S"),
        "count": len(files), "skipped": skipped, "by_ext": by_ext,
        "ocr_available": OCR_AVAILABLE, "files": files,
    }
    if callable(index_path):          # delegated saving (database)
        index_path(index)
    else:
        tmp = index_path + ".tmp"
        with io.open(tmp, "w", encoding="utf-8") as f:
            json.dump(index, f, ensure_ascii=False)
        os.replace(tmp, index_path)
    done = index_summary(index)
    done["phase"] = "done"
    yield done


def scan_folder(folder, index_path, base_dir=None):
    """Non-streaming version: scans, saves and returns the summary."""
    summary = None
    for ev in scan_folder_stream(folder, index_path, base_dir):
        if ev.get("phase") == "done":
            summary = {k: v for k, v in ev.items() if k != "phase"}
    return summary


def index_summary(index):
    out = {k: index.get(k) for k in
           ("folder", "scanned_at", "count", "skipped", "by_ext", "ocr_available")}
    out["folder_rel"] = index.get("folder_rel")
    out["portable"] = bool(index.get("folder_rel"))
    return out


def load_index(index_path):
    try:
        with io.open(index_path, "r", encoding="utf-8-sig") as f:
            return json.load(f)
    except Exception:
        return None


# ---------------------------------------------------------------- search ---
PHONE_QUERY = re.compile(r"^[\d\s.\-()+/]+$")   # query made only of digits and separators
_SEP = r"[\s.\-()/]*"                            # separators allowed between digits


_HAN = re.compile(r"[\u3400-\u4DBF\u4E00-\u9FFF\uF900-\uFAFF]")


def _has_han(s):
    """Chinese/Japanese ideographs: the text has no spaces between words, so
    'whole word' mode has no boundaries to look for and the sequence is searched."""
    return bool(_HAN.search(str(s or "")))


def build_matcher(nq, mode):
    """Build the search regex on the normalised text (1:1 with the original
    text). If the query is a NUMBER, the digits may be separated by
    spaces/hyphens/dots: so '+23562798882' also finds '+235 62 79 88 82'."""
    digits = re.sub(r"\D", "", nq)
    if len(digits) >= 5 and PHONE_QUERY.match(nq):
        core = _SEP.join(re.escape(d) for d in digits)
        pat = r"(?:\+" + _SEP + r")?" + core             # optional leading +
        if mode == "total":
            pat = r"(?<!\d)" + pat + r"(?!\d)"           # number boundaries
        return re.compile(pat)
    esc = re.escape(nq)
    if mode == "total" and not _has_han(nq):
        return re.compile(r"(?<!\w)" + esc + r"(?!\w)")
    return re.compile(esc)


def _snippets_from(text, matches, max_snips=3, radius=60):
    """Context excerpts around the matches. Replacements are 1:1
    (newline/tab -> space) so offsets stay valid; the browser collapses
    multiple spaces when rendering."""
    out = []
    for m in matches[:max_snips]:
        i, j = m.start(), m.end()
        a = max(0, i - radius)
        b = min(len(text), j + radius)
        pre = "…" if a > 0 else ""
        post = "…" if b < len(text) else ""
        snip = text[a:b].replace("\n", " ").replace("\r", " ").replace("\t", " ")
        out.append({"text": pre + snip + post,
                    "at": i - a + len(pre),      # offset of the match in the snippet
                    "len": j - i})               # REAL length of the match
    return out


def search_index(index, query, mode="partial", max_files=60):
    """Search 'query' in the index. mode: 'partial' (substring) or 'total'
    (whole word). Returns the files containing the term, with previews.
    For phone numbers separators are ignored."""
    nq = norm(query).strip()
    if not index or not nq:
        return {"query": query, "mode": mode, "count": 0, "results": []}

    matcher = build_matcher(nq, mode)
    results = []
    for f in index.get("files", []):
        text = f.get("text") or ""
        # use the already saved normalised text (fast, aligned 1:1); recompute it
        # only if missing altogether (old indexes without the 'ntext' field).
        nt = f.get("ntext")
        if nt is None:
            nt = norm(text)
        if not nt:
            continue
        matches = list(matcher.finditer(nt))
        if not matches:
            continue
        results.append({
            "id": f["id"], "name": f["name"], "rel": f["rel"],
            "ext": f["ext"], "occ": len(matches), "ocr": f.get("ocr", False),
            "snippets": _snippets_from(text, matches),
        })
    results.sort(key=lambda r: -r["occ"])
    total = len(results)
    return {"query": query, "mode": mode, "count": total,
            "results": results[:max_files], "truncated": total > max_files}


_WORDSEP = r"[\s.\-,/()]+"        # separator between the words of a phrase


def _term_core(w):
    """Regex fragment for one word: separator-tolerant if it is a number."""
    nw = norm(w)
    digits = re.sub(r"\D", "", nw)
    if len(digits) >= 5 and PHONE_QUERY.match(nw):
        return r"(?:\+" + _SEP + r")?" + _SEP.join(re.escape(d) for d in digits)
    return re.escape(nw)


def _combo_regex(words, mode):
    """Regex of a combination: the words must appear consecutively
    (separated by spaces/punctuation). For a single word, in 'total' mode
    it adds word boundaries."""
    core = _WORDSEP.join(_term_core(w) for w in words)
    if mode == "total" and len(words) == 1 and not _has_han(words[0]):
        core = r"(?<!\w)" + core + r"(?!\w)"
    return re.compile(core)


def search_multi(index, terms, mode="partial", max_files=80, max_words=4, max_combos=400):
    """Search the sources for ALL combinations of the words of the card key:
    from 1 to max_words words, in every order (permutations). Returns the
    combinations that produced results (for the chips) and the files that
    contain them. A per-word pre-filter keeps times low even on large
    indexes: a combination is searched only in files that contain all of
    its words."""
    seen, words = set(), []
    for t in terms:
        t = str(t or "").strip()
        k = norm(t)
        if len(t) >= 2 and k and k not in seen:
            seen.add(k)
            words.append(t)
    if not index or not words:
        return {"terms": [], "results": [], "total": 0, "truncated": False}

    files = index.get("files", [])
    nts = []
    for f in files:
        nt = f.get("ntext")
        nts.append(nt if nt is not None else norm(f.get("text") or ""))

    # 1) pre-filter: which files contain each word
    word_rx = {w: _combo_regex([w], "partial") for w in words}
    word_files = {w: set() for w in words}
    for i, nt in enumerate(nts):
        if not nt:
            continue
        for w in words:
            if word_rx[w].search(nt):
                word_files[w].add(i)

    # 2) generate the combinations (permutations) of 1..max_words words
    combos, seen_c = [], set()
    for r in range(1, min(max_words, len(words)) + 1):
        for perm in itertools.permutations(words, r):
            key = "\x00".join(norm(x) for x in perm)
            if key in seen_c:
                continue
            seen_c.add(key)
            combos.append(perm)
            if len(combos) >= max_combos:
                break
        if len(combos) >= max_combos:
            break

    # 3) search each combination as a PHRASE, only in the candidate files
    combo_count = {}
    file_data = {}
    for perm in combos:
        cand = set.intersection(*[word_files[w] for w in perm]) if perm else set()
        if not cand:
            continue
        label = " ".join(perm)
        rx = _combo_regex(list(perm), mode)
        cnt = 0
        for i in cand:
            ms = list(rx.finditer(nts[i]))
            if not ms:
                continue
            cnt += 1
            fd = file_data.setdefault(i, {"matched": [], "occ": 0, "snippets": []})
            fd["matched"].append(label)
            fd["occ"] += len(ms)
            for sn in _snippets_from(files[i].get("text") or "", ms, max_snips=1):
                sn["term"] = label
                fd["snippets"].append(sn)
        if cnt:
            combo_count[label] = cnt

    # chips: combinations with more words at the top, then by count
    terms_sorted = sorted(
        ({"term": l, "count": c, "words": l.count(" ") + 1}
         for l, c in combo_count.items()),
        key=lambda x: (-x["words"], -x["count"], x["term"]))

    results = []
    for i, fd in file_data.items():
        f = files[i]
        # snippets: longest (most specific) combinations first
        fd["snippets"].sort(key=lambda s: -(s["term"].count(" ") + 1))
        results.append({
            "id": f["id"], "name": f["name"], "rel": f["rel"], "ext": f["ext"],
            "ocr": f.get("ocr", False), "matched": fd["matched"],
            "occ": fd["occ"], "snippets": fd["snippets"][:8],
        })
    results.sort(key=lambda r: (-max((m.count(" ") + 1) for m in r["matched"]),
                                -len(r["matched"]), -r["occ"]))
    return {"terms": terms_sorted, "results": results[:max_files],
            "total": len(results), "truncated": len(results) > max_files}


def find_file(index, fid):
    for f in index.get("files", []):
        if f["id"] == fid:
            return f
    return None
